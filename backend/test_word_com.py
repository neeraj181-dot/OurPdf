import os
import tempfile
import win32com.client
import pythoncom
import docx

def test_word_com_conversion():
    # 1. Create a sample .docx or .doc
    doc = docx.Document()
    doc.add_heading("Native Word COM PDF Export Test", level=0)
    doc.add_paragraph("This is rendered using Microsoft Word native COM engine.")
    
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        docx_path = f.name
        doc.save(docx_path)
        
    pdf_path = docx_path.replace(".docx", ".pdf")
    
    print(f"Created docx at: {docx_path}")
    print(f"Target pdf at: {pdf_path}")
    
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    
    try:
        abs_docx = os.path.abspath(docx_path)
        abs_pdf = os.path.abspath(pdf_path)
        
        doc_obj = word.Documents.Open(abs_docx)
        # 17 = wdFormatPDF
        doc_obj.SaveAs(abs_pdf, FileFormat=17)
        doc_obj.Close(SaveChanges=False)
        print("Successfully exported PDF via Word.Application!")
        
        with open(abs_pdf, "rb") as pf:
            pdf_bytes = pf.read()
            print(f"Generated PDF bytes: {len(pdf_bytes)}")
    finally:
        word.Quit()
        pythoncom.CoUninitialize()
        if os.path.exists(docx_path):
            os.remove(docx_path)
        if os.path.exists(pdf_path):
            os.remove(pdf_path)

if __name__ == "__main__":
    test_word_com_conversion()
