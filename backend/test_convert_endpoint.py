import io
import docx
from fastapi.testclient import TestClient
from app.main import app
import pymupdf

client = TestClient(app)

def test_word_to_pdf_endpoint():
    # 1. Create a real sample DOCX in memory
    doc = docx.Document()
    doc.add_heading("API Word to PDF Test", level=0)
    p = doc.add_paragraph("This is a paragraph converted via FastAPI endpoint. ")
    p.add_run("Formatted in bold.").bold = True
    
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Col 1"
    t.cell(0, 1).text = "Col 2"
    t.cell(1, 0).text = "Val A"
    t.cell(1, 1).text = "Val B"
    
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()
    
    # 2. Call POST /api/convert/word-to-pdf
    response = client.post(
        "/api/convert/word-to-pdf",
        files={"file": ("sample_document.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    
    assert response.status_code == 200, f"Expected status 200, got {response.status_code}: {response.text}"
    assert response.headers["content-type"] == "application/pdf"
    
    pdf_bytes = response.content
    assert len(pdf_bytes) > 500, "PDF content too small"
    
    # 3. Verify PDF structure with PyMuPDF
    pdf_doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    assert len(pdf_doc) >= 1
    page_text = pdf_doc[0].get_text()
    print("Endpoint Extracted Text:")
    print(page_text.strip())
    assert "API Word to PDF Test" in page_text
    assert "Col 1" in page_text
    
    print("\nSUCCESS: /api/convert/word-to-pdf endpoint test passed 100%!")

if __name__ == "__main__":
    test_word_to_pdf_endpoint()
